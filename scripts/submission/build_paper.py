"""검토 가능한 분석표에서 공모전 논문과 정적 그림을 만든다."""
import json
from pathlib import Path
from scripts.submission.audit import ROOT, OUT, read_csv
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build():
    s=json.loads((OUT/'analysis_summary.json').read_text(encoding='utf-8'))
    daily=read_csv(OUT/'observation_daily.csv')
    gaps=read_csv(OUT/'facility_gap_candidates.csv')
    font=Path('C:/Windows/Fonts/malgun.ttf')
    if font.exists():
        font_manager.fontManager.addfont(str(font));plt.rcParams['font.family']='Malgun Gothic'
    plt.rcParams.update({'axes.unicode_minus':False,'font.size':11,'axes.spines.top':False,'axes.spines.right':False})
    fig,ax=plt.subplots(figsize=(7.1,2.4))
    eligible=[r for r in gaps if r['publication_status']=='candidate_requires_human_review']
    for i,r in enumerate(eligible):
        lo,hi=int(r['candidate_gap_min'])/365.25,int(r['candidate_gap_max'])/365.25
        ax.barh(i,lo,color='#355C7D',height=.46)
        ax.plot([lo,hi],[i,i],color='#D18F30',linewidth=5)
        ax.text(hi+.09,i,f"{r['candidate_gap_min']}–{r['candidate_gap_max']}일",va='center',fontsize=10)
    ax.set_yticks(range(len(eligible)),[r['development'] for r in eligible]);ax.invert_yaxis();ax.set_xlim(0,11.3);ax.set_xlabel('입주 후보월부터 특정 철도시설 개통까지 경과 연수');ax.grid(axis='x',alpha=.15);fig.tight_layout();fig.savefig(OUT/'figure_1_facility_intervals.png',dpi=220);plt.close(fig)
    fig,ax=plt.subplots(figsize=(7.1,2.6))
    x=list(range(len(daily)));valid=[int(r['with_items_slots']) for r in daily];empty=[int(r['empty_slots']) for r in daily];missing=[int(r['missing_slots']) for r in daily]
    ax.bar(x,valid,color='#355C7D',label='도착정보 있음');ax.bar(x,empty,bottom=valid,color='#D18F30',label='정상 빈 응답');ax.bar(x,missing,bottom=[a+b for a,b in zip(valid,empty)],color='#D8DEE5',label='성공 응답 없음')
    for i,r in enumerate(daily):ax.text(i,49,f"{r['successful_unique_slots']}/48",ha='center',fontsize=10)
    ax.set_xticks(x,[r['obs_date'][5:] for r in daily]);ax.set_ylim(0,56);ax.set_ylabel('명목 5분 슬롯 수');ax.legend(ncol=3,loc='upper center',bbox_to_anchor=(.5,1.22),frameon=False,fontsize=9);fig.tight_layout();fig.savefig(OUT/'figure_2_observation_coverage.png',dpi=220);plt.close(fig)
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    normal=d.styles['Normal'];normal.font.name='맑은 고딕';normal.font.size=Pt(10.5);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕');normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(5)
    for sn,size in [('Title',26),('Heading 1',17),('Heading 2',12)]:
        st=d.styles[sn];st.font.name='맑은 고딕';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st._element.rPr.rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    def p(t):d.add_paragraph(t)
    pending_page=[False]
    def h(t):
        par=d.add_heading(t,1);par.paragraph_format.page_break_before=pending_page[0];pending_page[0]=False
        par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(9)
    def sub(t):
        par=d.add_heading(t,2);par.paragraph_format.space_before=Pt(8);par.paragraph_format.space_after=Pt(5)
    def table(headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,x in zip(t.rows[0].cells,headers):c.text=x
        trpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for row in rows:
            cells=t.add_row().cells
            for c,x in zip(cells,row):
                c.text=str(x)
                for par in c.paragraphs:
                    par.paragraph_format.space_after=Pt(4)
                    for run in par.runs:run.font.size=Pt(9)
    def page():pending_page[0]=True
    d.add_paragraph('교통데이터 분석 공모전',style='Subtitle')
    d.add_heading('신도시 교통약속과\n실제 서비스의 증거 기반 감사',0)
    p('공식 문서의 날짜 불확실성과 버스 도착정보의 관측 한계를 중심으로')
    p('박도현  |  Transport Gap Audit  |  2026년 9월 12일')
    sub('초록')
    p(f"본 연구는 신도시 입주와 교통시설 공급의 시차를 감사할 때 발생하는 증거의 불일치를 분석한다. 과거 5개 지구와 전방 모니터링 5개 지구를 대상으로 사건 레지스트리 {s['registry_rows']}행, PDF {s['pdf_files']}개, 공동주택 {s['parquet_profiles']['kapt_basic']['rows']:,}행 및 건축물 표제부 {s['parquet_profiles']['bldg_title']['rows']:,}행을 점검하고, 남양주 한 정류소의 2026년 9월 4일부터 11일까지 저장된 도착정보를 재집계하였다. 활성 사건 후보 180행 중 현재 파일 기반 필수항목 검사는 107행(59.4%)이 통과하여 기존 요약 117행과 차이를 보였다. 사람 검증 완료 사건은 0행이었다. 세 지구의 입주 후보월과 특정 철도 개통일 간 계산 범위는 2,231~3,376일이지만, 이는 전체 대중교통 부재를 뜻하지 않는다. 버스 도착정보는 259개 성공 응답, 6,354개 예측 항목, 15개 노선을 포함했으며 차량 식별자가 없어 실제 배차 준수율을 확정할 수 없었다. 이에 시설·구간·목표의 의미를 맞춘 약속 이력관리와 입주 전 공식 시간표 확보를 제안한다.")
    p('주제어  신도시 교통정책, 공공데이터 감사, 날짜 구간, 증거 추적, 버스 도착예정정보')
    p('분석 코드 및 부속자료  https://github.com/lythean1004/transport-gap-audit')
    p('저장소는 비공개 연구자료 보관용이며 심사 공유 시 별도 접근 권한이 필요하다.')
    page();h('1 분석 배경 및 목적')
    p('입주와 교통시설 개통 사이의 시간차는 주민의 생활 여건을 검토하는 출발점이다. 그러나 철도 개통 이전에도 기존 철도·버스가 존재할 수 있고, 특정 사업의 준공 목표는 영업운행 개시 목표와 다르다. 이 차이를 무시하면 장기간의 시설 대기 시간을 대중교통의 전면적 부재나 확정된 약속 불이행으로 과장할 수 있다.')
    p('본 연구의 질문은 세 가지다. 첫째, 기존 수집자료와 계산 결과가 원문까지 재현되는가. 둘째, 같은 교통시설의 입주·목표·운행 날짜를 어디까지 비교할 수 있는가. 셋째, 실시간 도착예정정보는 서비스 수준의 어떤 부분을 관측하며 무엇을 입증하지 못하는가. 분석 단위는 문서 사건, 공동주택 단지, 건축물 표제부, 정류소별 명목 시간 슬롯으로 분리하였다.')
    sub('활용 데이터')
    table(['자료와 제공기관','저장 규모','출처 및 선정 이유'],[
        ['사건 레지스트리와 공식 문서\n국토교통부·LH·GH·지자체','197행 / PDF 112개\n고유 PDF 92개','각 기관 원문 URL [1–6]\n입주·목표·개통의 근거 추적'],
        ['공동주택 기본정보\nK-apt','21,690행','공동주택관리정보시스템 [7]\n단지·세대수·사용승인 후보'],
        ['건축물 표제부\n국토교통부 건축HUB','4,067행','공공데이터포털 [8]\n공동주택 사용승인일 교차검토'],
        ['TAGO 정류소·노선정보\n국토교통부','전체 정류소 8,420행\n주변 조회 20,254행\n노선 서비스 12행','공공데이터포털 [9–11]\n좌표·정류소·공표 운영정보'],
        ['TAGO 도착예정정보\n국토교통부','원장 262행\n성공 JSON 259개','공공데이터포털 [10]\n2026-09-04~09-11 관측 감사'],
    ])
    p('과거 지구는 화성 동탄2·김포 한강·위례·남양주 다산·하남 미사이며, 전방 지구는 인천 계양·남양주 왕숙·하남 교산·고양 창릉·부천 대장이다. 정류소 관측은 남양주의 가운휴먼시아·다산효성해링턴타워 인근 1개 정류소에 한정된다. 이 표본을 10개 지구의 대표 표본으로 취급하지 않았다.')
    p('아이디어 180개 재평가 문서와 ACCESS-LINK 초안은 연구 주제 변경의 배경자료다. 레지스트리의 활성 사건 180행과는 별개의 자료이며, 초기 초안의 이동체인 성공확률·예산효과는 실제 산출 결과로 사용하지 않았다.')
    page();h('2 분석과정 및 방법')
    sub('원본 보존과 증거 추적')
    p('원본은 변경하지 않고 SHA-256과 경로를 목록화하였다. PDF 112개를 기존 파일 목록과 대조했으며 모두 같은 해시를 유지했다. 같은 바이트의 PDF 20개는 고유 증거 수에서 중복 계산하지 않았다. 출처 URL·파일 존재·해시·페이지·발췌·날짜·수단·노선 등 필수항목을 현재 자료로 다시 검사했다. 파일 보존, 기계적 필드 검사, 의미 검토, 사람 검증은 서로 다른 검증 단계로 관리하였다.')
    sub('날짜 정밀도와 비교 가능성')
    p('월 단위 사건은 해당 월의 첫날과 마지막 날, 연 단위는 1월 1일과 12월 31일, 반기는 해당 반기의 시작과 끝을 경계로 둔다. 입주 후보 구간이 [입주하한, 입주상한]이고 개통 구간이 [개통하한, 개통상한]이면 경과기간의 하한은 개통하한−입주상한, 상한은 개통상한−입주하한이다. 이는 날짜 정밀도에 따른 가능한 범위이며 통계적 신뢰구간이 아니다.')
    p('약속 비교는 같은 수단·시설·구간뿐 아니라 목표 동사까지 확인한다. 준공과 영업개시는 분리한다. 특정 기간 안의 개통 목표는 기간 종료를 기준으로 지연 여부를 해석하되, 구간 전체와의 산술 차이도 부속표에 남긴다. 개통 근거가 없다는 이유만으로 미개통 또는 미이행을 확정하지 않는다. 단지 사용승인일도 실제 첫 입주일과 같다고 가정하지 않는다.')
    sub('도착정보와 결측 처리')
    p('수집 설계는 오전 07:00~08:55와 오후 17:00~18:55, 5분 간격으로 하루 48개 명목 슬롯이다. 실제 수집된 6일을 모두 보고하고 최초 셰이크다운일과 이후 5일을 구분한다. 같은 슬롯의 재시도는 성공 슬롯 수에 중복 산입하지 않는다. 실제 호출시각이 120초 이내인 인접 호출은 별도 민감도 항목으로 표시한다.')
    p('resultCode가 정확히 00이면 항목 존재 여부에 따라 정상 유항목·정상 빈 응답으로 구분하고 그 외는 API 오류로 분류한다. 빈 응답은 해당 응답에 운행 중 버스 항목이 없다는 뜻으로 기록하며, 정류소 무효나 역사적 서비스 부재를 뜻하지 않는다. arrtime은 초 단위 도착예정시간이고 같은 차량이 반복 나타날 수 있다. 차량 식별자가 없으므로 예측 항목 수를 운행 횟수나 탑승객 수로 세지 않는다.')
    sub('분석도구와 재현 범위')
    p('Python, pandas·Polars, PyArrow, PyMuPDF, pytest로 파일·구조·계산을 검사하고 Matplotlib으로 그림을 작성했다. 이번 재분석의 추가 공공 API 호출은 0회다. 과거 수집을 새로 실행하지 않고 보관된 응답을 재파싱했다. 코드의 실제 네트워크 운영 안정성과 장기 관측의 대표성은 이번 오프라인 검사 범위 밖이다.')
    page();h('3 공식 문서와 과거 사건 분석 결과')
    p('활성 후보 180행 중 필수항목 통과는 107행(59.4%), 미통과는 73행(40.6%)이었다. 기존 저장 요약 117행에서 10행이 추가로 미통과했다. 주된 문제는 직접 출처 URL 부재 30행, 사건일 부재 26행, 발췌문 불일치 22행이다. 한 행에 여러 문제가 있을 수 있어 건수를 합산하면 안 된다. 사람 검증 완료 사건은 0행이며 기존 엄격 게이트의 과거 지구 5곳 FAIL 상태를 유지한다.')
    table(['지구와 시설','입주 후보 → 개통','계산 범위 및 처리'],[
        ['동탄2 / GTX-A 수서~동탄','2015년 1월 → 2024-03-30','3,346~3,376일 / 후보'],
        ['김포 한강 / 김포골드라인','2011년 6월 → 2019-09-28','3,012~3,041일 / 후보'],
        ['하남 미사 / 하남선 1단계','2014년 6월 → 2020-08-08','2,231~2,260일 / 후보'],
        ['다산 / 별내선','2018년 본격 입주 문구','최초 입주 입증 부족 / 제외'],
        ['위례 / 위례선','2013년 12월 입주 후보','현재 개통 여부 입증 부족 / 제외'],
    ])
    d.add_picture(str(OUT/'figure_1_facility_intervals.png'),width=Inches(6.65))
    p('그림 1. 세 지구의 입주 후보월부터 특정 철도 개통까지의 계산 범위. 주황선은 날짜 구간의 폭이다. 사람 검증 전 후보이며 대중교통 무서비스 기간이 아니다. 출처: 사건 레지스트리의 SRC-H001·H005·H007·H011·H025·H029, facility_gap_candidates.csv.')
    p('GTX-A의 2024년 상반기 목표와 3월 30일 개통은 목표 기간 안에 있다. 기존 −92일은 반기 말 기준 차이이므로 정확히 92일의 조기 이행으로 단정하지 않는다. 김포의 2018년 목표와 2019년 9월 28일 개통 간 산술 차이는 271~635일이며 목표 연말 초과분은 271일이다. 미사는 선택된 2020년 8월 8일 목표와 개통일이 같다. 모두 선택한 약속 버전에 대한 비교이므로 최초 약속 대비 사업 전체 지연을 의미하지 않는다.')
    p('다산 자료의 “2022년 준공 예정”을 2024년 영업개시와 비교해 588일의 개통 지연으로 표시한 결과는 본안에서 제외하였다. 전방 5개 지구도 예정 블록과 노선 목표의 의미 검증이 남아 있어 미래 미이행 순위를 산출하지 않았다.')
    page();h('4 현재 관측과 교차검토 결과')
    p('원장 262행은 정상 유항목 254회, 정상 빈 응답 5회, API 오류 3회로 구성된다. 성공한 JSON 259개에서 6,354개 예측 항목과 15개 노선을 확인했으며 원장의 항목 수와 원 응답 항목 수는 일치했다. 이는 6,354대의 버스가 도착했다는 의미가 아니다.')
    d.add_picture(str(OUT/'figure_2_observation_coverage.png'),width=Inches(6.65))
    p('그림 2. 날짜별 성공 슬롯과 빈 응답·미확보 슬롯. 각 날짜 분모는 예정 48슬롯이며 재시도 실패 행을 별도 슬롯으로 세지 않았다. 9월 4일은 셰이크다운일이다. 출처: arrival_ledger.csv 및 원 JSON, observation_daily.csv.')
    table(['관측 범위','성공 / 계획 슬롯','해석'],[
        ['6일 전체','259 / 288 (89.9%)','성공 응답 없는 슬롯 29개'],
        ['9월 7~11일','229 / 240 (95.4%)','정규 5일, 미확보 11개'],
        ['9월 8일 근접 호출 민감도','48 → 47개','120초 이내 인접 호출 1쌍'],
        ['기존 노선정보 조회 범위','12 / 관측 15개 노선','추가 3개 노선 공표정보 비교 불가'],
    ])
    p('실패 요청 3회 중 2회는 동일 슬롯의 성공 응답으로 복구되었고, 나머지는 해당 슬롯 성공 응답이 없다. 원장 자체가 없는 28슬롯과 합쳐 미확보 슬롯은 29개다. 9월 8일 근접 호출 한 쌍을 하나로 취급하면 독립 시점으로 보기 어려운 관측 1개가 줄어들어 전체는 258/288, 정규 5일은 228/240이 된다. 이는 샘플링 품질의 민감도이며 서비스 결행률이 아니다.')
    p('기존 headway_observed 72행은 12개 노선×3일×오전·오후의 과거 가공결과다. 현재 저장소의 정식 소스 경로에서 생성 코드를 확인하지 못했고 차량 추적도 불가능해 이를 실측 배차로 승격하지 않았다. 노선별 간격의 중앙값을 정류소 통합 서비스 빈도로 읽는 것도 타당하지 않다. 이번 논문에서는 재현 가능한 응답·슬롯 품질을 주된 관측 결과로 제시한다.')
    p('단지별 사용승인 교차표 289행은 일치 181행(62.6%), 30일 이내 차이 1행, 충돌 1행, 건축HUB 대응값 부재 106행(36.7%)이다. 대응값이 있는 183행만 보면 일치는 98.9%이므로 분모를 명시해야 한다. 사용승인과 첫 입주는 다르고 입주 세대의 시계열도 없어 세대·일 노출량은 계산하지 않았다.')
    page();h('5 정책제안 및 기대효과')
    sub('약속의 시설과 목표를 분리한 공개 이력관리')
    p('사업시행자와 지자체는 수단·노선·구간·준공 또는 영업개시 목표·발표시점·변경 사유를 한 행의 약속 버전으로 공개할 필요가 있다. 최근 목표로 과거 약속을 덮어쓰지 않고 변경 이력을 유지하면 주민은 일정 변경을 추적할 수 있고, 행정기관은 어떤 목표에 대한 평가인지 설명할 수 있다. 우선 적용 대상은 본 연구의 과거 5개 지구에서 확인된 약속 버전이다.')
    sub('입주 전 최소서비스의 문서 확보')
    p('입주 지원을 담당하는 지자체와 버스 인가·운영기관은 입주 블록별 정류소, 방향별 시간표, 첫차·막차, 적용 시작일을 함께 공개하는 방안을 검토할 수 있다. 전방 5개 지구에는 현재 미이행 딱지를 붙이는 대신 공식 시간표와 운행개시 증빙의 확보 현황을 관리한다. 500·800·1,000m 및 15·20·30분은 기존 프로젝트의 탐색용 임계값이며 법정 기준이나 이미 입증된 적정 수준으로 제시하지 않는다.')
    sub('정보 품질과 서비스 성과를 분리한 운영지표')
    p('API 성공 슬롯 비율, 원문·발췌 일치율, 사람 검증률은 자료의 품질지표로 관리하고, 실제 배차·통행시간·접근 가능 목적지 수는 별도 현장 또는 운행자료로 평가한다. 기대효과는 과장된 지연 판정의 감소와 증거 확인 비용의 절감이다. 효과의 크기, 예산 대비 편익 또는 주민 통행 개선량은 이번 자료로 추정하지 않았다.')
    sub('AI 기술 활용')
    p('기존 Antigravity 작업은 자료 발견·수집 코드·정리·1차 분석에 활용되었다. 이번 Codex 작업은 코드 및 보관자료 감사, 회귀 테스트, 구간 계산, 표·그림·논문 작성과 배포용 비밀정보 검사에 활용했다. 주요 지시는 “공식 근거와 결과를 연결하라”, “관측 스냅샷만으로 서비스 충족을 판정하지 말라”, “하드코딩과 키 노출을 점검하라”였다. 첨부 프롬프트는 참고자료로만 읽었다. AI 추출 정확도는 사람 골드셋이 없어 수치화하지 않았다.')
    sub('한계와 후속 검증')
    p('사람 검증 전 사건을 확정 사실로 승격하지 않았다는 점이 가장 큰 제한이다. 일부 입주 후보는 제목이나 회고적 설명에 의존하며 독립 자료의 의미 대조가 필요하다. 관측은 1개 정류소·6일·출퇴근 창에 한정되고 수집 누락이 무작위라는 근거도 없다. 차량 ID와 정류소별 공식 시간표가 없어 실제 배차 준수율, 기본서비스 충족 여부, 과거 버스 부재를 판정할 수 없다. 본 결과는 신도시 전체의 교통 불평등에 대한 인과 추정이 아니라 증거 감사와 제한된 사례 분석이다.')
    p('제출 전 핵심 검수 대상은 세 지구의 입주·개통 원문, 약속 목표의 버전 및 의미, 발췌 불일치 22건이다. 공개 배포 시에는 원문별 이용조건을 확인하고 신청서·서명·재학증명서·인증키를 연구자료와 분리해야 한다.')
    page();h('참고문헌 및 데이터 출처')
    p('아래 기관 원문과 제공 자료를 사용했다. 개별 사건의 원 URL·파일 경로·해시는 동봉한 source_registry.csv 및 file_inventory.csv에서 확인할 수 있다. 기관 간 원문을 별도의 독립 조사로 오인하지 않도록 동일 문서의 사본은 구분하였다.')
    refs=[
        '[1] LH. 수도권 최대 신도시 동탄(2)지구 첫 입주 시작. 2015-01-23. https://www.lh.or.kr/gallery.es?act=view&bid=0003&list_no=7977&mid=a10502000000',
        '[2] 국토교통부. 수원호매실·화성동탄2 광역교통 특별대책. 2022-10-26. https://www.molit.go.kr/mta/USR/N0201/m_36770/dtl.jsp?id=95087345',
        '[3] 국토교통부. GTX-A 수서~동탄 3월 30일 개통. 2024-03-29. https://www.molit.go.kr/USR/NEWS/m_72/dtl.jsp?id=95089607',
        '[4] 김포시. 김포 지하철시대 열렸다 28일 첫 열차 출발. 2019-09. https://gimpo.go.kr/news/selectBbsNttView.do?bbsNo=466&key=9377&nttNo=94282',
        '[5] 경기도. 다산신도시 입주지원 특별대책반 본격 가동. 2018-02-12. https://gnews.gg.go.kr/briefing/brief_gongbo_view.do?BS_CODE=S017&number=36361',
        '[6] 하남시. 하남선 1단계 개통. 2020-08. https://www.hanam.go.kr/sosik/selectBbsNttView.do?bbsNo=1164&key=10048&nttNo=210883',
        '[7] 한국부동산원. 공동주택관리정보시스템 K-apt. https://www.k-apt.go.kr/ (보관 파일의 추출시각·해시를 데이터에 수록)',
        '[8] 국토교통부. 건축HUB 건축물대장정보 서비스. 공공데이터포털. https://www.data.go.kr/data/15134735/openapi.do',
        '[9] 국토교통부. TAGO 버스정류소정보. 공공데이터포털. https://www.data.go.kr/data/15098534/openapi.do',
        '[10] 국토교통부. TAGO 버스도착정보. 공공데이터포털. https://www.data.go.kr/data/15098530/openapi.do',
        '[11] 국토교통부. TAGO 버스노선정보 활용가이드. 2025-06-16. getRouteInfoIem, intervaltime 등. 저장 문서: docs/tago_route_api_findings.md 및 원 활용가이드.',
        '[12] 숲과나눔. AI와 함께하는 교통문제 해결을 위한 데이터 분석 공모전 공고 및 붙임 분석보고서 양식. 2026-08-18. https://koreashe.org/notice/?mod=document&uid=95427',
    ]
    for r in refs:
        par=d.add_paragraph(r);par.paragraph_format.space_after=Pt(8)
        for run in par.runs:run.font.size=Pt(9)
    p('부속자료: 전체 데이터 프로파일, 코드 감사, 발췌 재검산, 원장 재파싱 표, 날짜 구간 표, 원본 및 배포본 해시 목록, 테스트 기록, 그림 원본, 재현 코드. 기존 시점 결과는 보관자료로 구분하며 이번 본문 수치와 혼용하지 않는다.')
    for element in [d._element,d.styles.element]:
        for border in element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    d.core_properties.title='신도시 교통약속과 실제 서비스의 증거 기반 감사';d.core_properties.author='박도현';d.core_properties.keywords='transport gap; evidence audit'
    d.save(OUT/'교통데이터공모전_논문_1차.docx')
    print('Created paper and 2 figures')


if __name__=='__main__':build()
