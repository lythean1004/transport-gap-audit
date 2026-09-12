"""V02 검수 이력·원문 날짜 재대조·관측 재현을 반영한 공모전 논문."""
import json
from pathlib import Path
from scripts.submission.audit import ROOT,read_csv
from scripts.submission.review_v02 import OUT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build():
    s=json.loads((OUT/'analysis_summary_v02.json').read_text(encoding='utf-8'));r=s['review_summary']
    daily=read_csv(OUT/'observation_daily.csv');gaps=read_csv(OUT/'facility_document_comparison_v02.csv')
    refs={x['source_id']:x for x in read_csv(OUT/'source_registry_v02.csv')}
    font_manager.fontManager.addfont('C:/Windows/Fonts/malgun.ttf');plt.rcParams.update({'font.family':'Malgun Gothic','axes.unicode_minus':False,'font.size':10,'axes.spines.top':False,'axes.spines.right':False})
    plot=[x for x in gaps if x['status']=='document_announced_start_comparison']
    fig,ax=plt.subplots(figsize=(7,2.3))
    for i,x in enumerate(plot):
        lo,hi=int(x['document_gap_min']),int(x['document_gap_max']);ax.barh(i,lo,color='#355c7d',height=.45);ax.plot([lo,hi],[i,i],lw=5,color='#c68c33');ax.text(hi+40,i,str(lo) if lo==hi else f'{lo}–{hi}',va='center')
    ax.set_yticks(range(len(plot)),[x['development'] for x in plot]);ax.invert_yaxis();ax.set_xlim(0,3800);ax.set_xlabel('공식 입주개시 발표의 기준일부터 선택 철도시설 개통일까지(일)');ax.grid(axis='x',alpha=.15);fig.tight_layout();fig.savefig(OUT/'figure_1_document_intervals_v02.png',dpi=220);plt.close(fig)
    fig,ax=plt.subplots(figsize=(7,2.4));x=range(len(daily));v=[int(z['with_items_slots']) for z in daily];e=[int(z['empty_slots']) for z in daily];m=[int(z['missing_slots']) for z in daily]
    ax.bar(x,v,label='도착정보 있음',color='#355c7d');ax.bar(x,e,bottom=v,label='정상 빈 응답',color='#c68c33');ax.bar(x,m,bottom=[a+b for a,b in zip(v,e)],label='성공 응답 없음',color='#d8dee5');ax.set_xticks(list(x),[z['obs_date'][5:] for z in daily]);ax.set_ylim(0,53);ax.set_ylabel('명목 5분 슬롯 수');ax.legend(ncol=3,loc='upper center',bbox_to_anchor=(.5,1.21),frameon=False);fig.tight_layout();fig.savefig(OUT/'figure_2_coverage_v02.png',dpi=220);plt.close(fig)
    d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    for name,size in [('Normal',10.5),('Title',24),('Heading 1',17),('Heading 2',12)]:
        st=d.styles[name];st.font.name='맑은 고딕';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
    st=d.styles['Normal'];st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(6)
    def p(t):d.add_paragraph(t)
    def h(t,new=False):
        q=d.add_heading(t,1);q.paragraph_format.page_break_before=new;q.paragraph_format.space_before=Pt(0)
    def sub(t):d.add_heading(t,2)
    def table(headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for c,x in zip(t.rows[0].cells,headers):c.text=x
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows:
            for c,x in zip(t.add_row().cells,row):
                c.text=str(x)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(4)
                    for run in para.runs:run.font.size=Pt(9)
    d.add_paragraph('교통데이터 분석 공모전 · 연구논문 V02',style='Subtitle')
    d.add_heading('신도시 입주와 교통시설 공급의\n시차에 대한 증거 기반 감사',0)
    p('사람 검수 이력, 공식 문서의 날짜 해석 및 도착예정정보 재현성을 중심으로')
    p('박도현 | Transport Gap Audit | 2026년 9월 12일')
    sub('초록')
    p(f"본 연구는 신도시 교통시설 공급 시차를 측정할 때 발생하는 날짜·자료계보·검증상태의 불일치를 분석한다. 과거 5개 지구와 전방 모니터링 5개 지구의 사건 레지스트리 {s['registry_rows']}행을 사용자 검수대장 및 최종 보완대장과 결합하였다. 사람 검수 기록은 {r['review_status_counts']['human_verified']}행이며, 원문 해시·검수자·검수 시각을 연결했다. 원문 재대조에서는 동탄2의 보도자료 배포일과 입주개시 예정일이 혼동된 사례, 김포 한강의 월 단위 공식 근거가 일 단위 값으로 정밀화된 사례를 확인하였다. 원문이 밝힌 입주개시 기준일과 선택 철도시설 개통일의 차이는 동탄2 3,347일, 김포 한강 3,012~3,041일, 하남 미사 2,231일이다. 이는 발표 기준일의 비교이며 전체 대중교통 부재나 실제 최초 입주일의 사후 검증을 의미하지 않는다. 별도로 한 정류소의 6일간 저장 관측 {s['raw_json_files']}개와 예측 항목 {s['parsed_items']:,}개를 재집계하고, 기존 배차 추정표 72행의 생성 경로를 재현하였다. 결과는 사람 검수 사실과 지표의 증거 충족 여부를 분리하고, 교통 약속의 시설·구간·목표시점 및 원문 위치를 함께 관리할 필요성을 보여준다.")
    p('주제어: 신도시 교통정책, 공식 문서, 자료계보, 사람 검수, 날짜 불확실성, 도착예정정보')
    sub('판본 안내')
    p('V01의 검수대장 연결 누락과 생성 코드 탐색 누락을 정정한 2차 초안이다. 검수 이력은 보존하되 원문 충돌을 별도로 기록한다. 증거 문구·페이지의 추가 정리가 필요한 항목과 실제 서비스 이행 미검증 영역은 본문에서 구분한다.')

    h('1. 연구 문제와 자료 범위',True)
    p('입주 시점보다 교통시설 공급이 늦어지면 주민은 다른 수단에 의존하거나 통행시간 증가를 경험할 수 있다. 그러나 특정 철도 개통이 늦었다는 사실만으로 그 기간 전체에 대중교통이 없었다고 판단할 수는 없다. 입주 날짜와 개통 날짜의 정의가 다르거나, 계획 확정일을 개통 약속으로 읽는 경우 측정치 자체도 달라진다. 본 연구는 피해 규모를 성급히 추정하기보다, 공모전에 제출할 근거와 계산을 추적 가능한 형태로 구성하는 데 목적을 둔다.')
    p('연구 질문은 세 가지다. 첫째, 사람 검수 결과가 실제 분석 입력에 연결되었는가? 둘째, 공식 원문의 날짜 정밀도와 사건 의미가 계산 과정에서 보존되는가? 셋째, 도착예정정보의 저장 관측과 기존 추정치가 재현되며 어떤 해석까지 허용되는가?')
    table(['자료','규모','분석상 용도와 한계'],[
        ['사건·검수대장',f"197행 / 검수 {r['review_status_counts']['human_verified']}행",'입주·약속·개통·기본서비스 사건의 근거 연결'],
        ['PDF 처리대장',f"{s['pdf_files']}개 / 고유 해시 {s['unique_pdf_hashes']}개",'문서 중복과 처리 이력 확인; HWP·HWPX 별도 포함'],
        ['공동주택 / 건축물',f"{s['parquet_profiles']['kapt_basic']['rows']:,} / {s['parquet_profiles']['bldg_title']['rows']:,}행",'단지·사용승인 정보; 실제 전입이나 최초 입주와 동일하지 않음'],
        ['도착정보',f"{s['raw_json_files']}개 JSON / {s['parsed_items']:,}항목",'한 정류소 6일, 15개 노선의 저장 예측정보'],
        ['지오코딩 보정','6건 / 작성자 AI','사람 검수 사건에 합산하지 않음'],
        ['정류소 검수','1건 / 지도 이미지 4개','검수자·시각 존재; 네트워크 전체를 대표하지 않음']])
    p('프로젝트와 제공된 다운로드 폴더의 파일을 목록화하고 중복 사본을 해시로 대조했다. 실행환경·Git 내부·테스트 캐시는 연구 근거와 구분했다. 다운로드 폴더의 V01 사본 3,991개는 배포본과 일치했다. 파일 판독 및 자동 검색은 원문 의미를 모두 검증했다는 뜻이 아니며, 이미지 및 문자열 불일치 항목은 별도 대장에 남겼다.')
    p('본 판본에서는 추가 API를 호출하지 않았다. 기준 데이터는 저장된 파일이며 미래 계획의 현행성이나 2026년 9월 12일 이후의 이행 여부를 주장하지 않는다.')

    h('2. 검수 결합과 측정 방법',True)
    p('검수대장에는 후보 식별자가 재사용된 행이 있어 후보 ID만으로 결합할 수 없다. 후보 ID, 원문 경로, SHA-256을 함께 사용하고 다중 후보는 지구명으로 구분했다. 그 결과 197행이 각각 하나의 원자료 행에 연결되었다. 더 상세한 최종 보완대장은 source_id·원문 해시·검수상태의 일치를 확인한 뒤 검수자, 검수 시각, 인용문, 페이지, 모순 사유를 결합했다. 파일명에 “최종”이 있다는 이유만으로 값을 채택하지 않았다.')
    table(['단계','판정 기준','보존 내용'],[['사람 검수','사용자 대장의 검수 상태','검수자·검수 시각·입력 파일 해시'],['원문 연결','같은 원문 버전 및 인용 위치','페이지·첨부파일·문구 대조'],['지표 비교','사건 의미·시설·날짜 범위 정합','충돌 및 보류 사유'],['출판 판단','계산 가능성과 사실 단정 구분','계획/관측/미확인을 본문에 표시']])
    p('날짜를 구간 I=[I하한,I상한], 개통일을 O=[O하한,O상한]으로 표현하면 경과기간은 [O하한−I상한, O상한−I하한]이다. 월만 제시한 자료를 임의의 특정 일로 정밀화하지 않는다. 연도 또는 반기 단위 약속도 같은 방식으로 다룬다. 연말을 기한으로 해석한 차이는 별도 “기한 대비 차이”이며, 해당 연도 전체 날짜 불확실성과 혼합하지 않는다.')
    p('원문 재대조 주석은 별도 CSV로 관리하고 원문 해시와 인용 문구가 일치할 때만 비교표에 적용했다. 사용자가 검수한 값은 삭제하지 않고 원문 기준값과 나란히 남겼다. 특히 동탄2는 예정 입주개시 발표, 김포는 월 단위 입주개시 설명, 미사는 당일 입주개시 발표로 구분한다. 이러한 자료만으로 각 가구의 실제 전입일을 확정하지 않는다.')
    p('도착정보는 resultCode=00인 정상 응답을 집계하며, items가 빈 정상 응답도 성공 응답으로 보존한다. 정상 빈 응답은 정류소 무효나 서비스 충족의 증거가 아니다. 차량 식별자가 없으므로 도착예측값의 증가·소실을 이용한 기존 이벤트 추정은 실제 차량 통과 관측과 구별한다.')

    h('3. 역사적 시차와 검수 결과',True)
    table(['지구','검수대장의 입주일','원문 재대조 결과','선택 시설까지'],[
        ['동탄2','2015-01-23','배포일 / 입주개시 예정 01-30','3,347일¹'],['김포 한강','2011-06-17','확보 공식 근거는 2011년 6월','3,012~3,041일¹'],['하남 미사','2014-06-30','당일 입주개시 발표','2,231일¹'],['위례','2013년 12월','선택 트램의 실제 개통 증거 미연결','보류'],['남양주 다산','2018년','최초 입주와 대책 발표 혼동 기록','보류']])
    d.add_picture(str(OUT/'figure_1_document_intervals_v02.png'),width=Inches(6.65))
    p('그림 1. 공식 발표의 입주개시 기준일과 선택 철도 개통일의 차이. ¹ 발표 기준 비교로서 실제 최초 입주의 사후 확정값이나 전체 대중교통 무서비스 기간이 아니다. 동탄2는 GTX-A 수서~동탄, 김포는 김포골드라인, 미사는 하남선 1단계를 비교했다. 각 기준일과 원문 해시는 document_date_checks_v02.csv에 기록했다.')
    p('동탄2의 2015년 LH 보도자료는 입주 초기 버스 21개 노선을 마련했다고도 설명한다. 따라서 GTX-A 개통까지의 긴 시차를 “교통수단이 전혀 없던 기간”으로 표현하면 원문의 취지와 충돌한다. 대체 교통수단의 존재와 최소 서비스 수준의 충족은 다시 별개의 질문이다.')
    p('약속 날짜의 정밀도에 따라 동탄2의 2024년 상반기 목표 대비 차이는 −92~89일, 김포의 2018년 목표 대비 차이는 271~635일, 미사의 2020년 8월 8일 목표 대비 차이는 0일이다. 동탄2는 상반기 말 기한 이전에 개통했지만 범위 전체를 단일한 “92일 조기개통”으로 치환하지 않았다. 다산은 준공 목표와 운행 개시를 혼합할 수 있어 지연 비교를 보류했다.')
    p(f"사람 검수 {r['review_status_counts']['human_verified']}행은 고유 원문 해시 {r['human_verified_unique_document_hashes']}개에 연결된다. 이 중 현재 문자열·필드 검사는 {r['human_verified_structural_pass']}행이 통과하고 {r['human_verified_structural_issues']}행에 보완 표시가 남는다. 요약·재서술된 문구가 exact_excerpt 열에 들어간 경우도 있으므로 문자열 불일치가 곧 검수 오류나 허위 자료를 뜻하지는 않는다. 미검수 15행과 검수 완료 후 증거 필드 보완 대상은 구분한다.")

    h('4. 저장 관측 재집계와 추정치 재현',True)
    success=sum(int(x['successful_unique_slots']) for x in daily);planned=sum(int(x['planned_slots']) for x in daily)
    p(f"2026년 9월 4·7·8·9·10·11일, 오전 07:00~08:55와 오후 17:00~18:55의 명목 5분 슬롯 {planned}개 중 성공 응답은 {success}개({success/planned:.1%})다. 정상 빈 응답 5개를 포함하며 도착예측 항목은 6,354개다. 9월 4일 준비 관측을 제외한 5일은 229/240개(95.4%)다. 일별 관측 누락을 숨기지 않고 성공 슬롯의 분모를 명시했다.")
    p('이 자료는 07~09시·17~19시에 배치로 실제 수집한 API 응답이다. 추가 제공된 run_arrival_snapshot.bat의 실행 경로를 확인했고, 일반·재시도 로그 262건은 원장 262행과 일치했다. 오전 120개·오후 139개 원파일의 항목 수 불일치는 없었다. 9월 8일 재시도 1건은 명목 슬롯보다 5분 이상 늦게 호출되어 실제 시각도 보존한다. 원자료와 계산된 배차 추정치는 구별한다.')
    d.add_picture(str(OUT/'figure_2_coverage_v02.png'),width=Inches(6.65))
    p('그림 2. 일별 수집 충족도. 정상 빈 응답은 성공 응답에 포함한다. 호출 시각이 120초 이내로 인접한 슬롯 한 쌍을 하나로 처리하면 전체 실효 슬롯은 258/288개다. 이는 수집 중복 민감도이며 버스 운행량의 감소가 아니다.')
    p('루트 외 scratch 폴더에서 배차 추정 코드의 여러 판본을 확인했다. 최신 보관 스크립트의 계산 함수를 외부 호출과 원자료 덮어쓰기 없이 실행하여 기존 72행·11열을 재현했다. 결측 표현(NaN과 pd.NA)을 통일하면 모든 열이 일치한다. 생성 경로는 확인됐지만 추정 방식 자체가 실제 배차를 식별한다는 증명은 아니다.')
    table(['선택 / 결측 유지','이벤트가 검출된 그룹','4개 미만 이벤트 그룹'],[['기존 3일·공표정보 12노선','72','26'],['전체 6일·공표정보 12노선','143','71'],['전체 6일·관측 15노선','170','97']])
    p('기존 코드는 분석일 3일과 공표정보가 없는 노선 3개 제외를 고정했다. 전체 관측일·노선으로 확장해도 풀링한 예측변화 간격의 중앙값은 20분이지만, 이벤트가 적은 그룹의 비중과 분모는 달라진다. 빈 응답을 0으로 처리한 반사실 민감도에서는 전체 15노선의 간격 하위사분위수가 15분에서 10분으로 이동했다. 이 민감도는 원자료를 채우지 않은 별도 실험이며 실제 배차 20분 또는 배차 준수율로 보고하지 않는다.')

    h('5. 정책 활용, 재현성 및 한계',True)
    p('정책적으로는 입주 승인과 함께 교통 약속의 대상 시설·구간·수단·목표 날짜·약속 판본을 등록하고, 변경된 목표를 이전 판본과 연결하는 방식이 필요하다. 보도자료 등록일, 사업계획 승인일, 준공 목표일, 개통일을 각각 다른 사건으로 저장하면 이번 감사에서 발견된 혼동을 줄일 수 있다. 입주 전 버스 등 대체 수단은 노선 존재뿐 아니라 공식 시간표와 적용일을 확보해 검토해야 한다.')
    p('전방 5개 지구는 실제 이행 성적표가 아니라 계획의 시점 관계를 모니터링하는 대상으로 유지한다. 인천 계양·고양 창릉·부천 대장은 선택 교통시설의 목표 날짜가 결측이다. 남양주 왕숙의 연도 단위 목표와 입주 예정월은 겹칠 수 있으므로 연말을 단정해 “지연 확정”으로 분류하지 않는다. 하남 교산은 확보 계획상 목표가 입주보다 앞서지만 실제 운행이나 안전한 통학·통근을 보장하는 결과로 해석하지 않는다.')
    p('자료별 검증 상태는 독립적이다. 사람 검수는 원문 검토의 기록이며, 원문 해시는 파일 버전의 동일성을 확인한다. 코드 재현은 같은 입력으로 같은 출력이 나오는지 확인한다. 어느 하나도 나머지의 타당성을 대신하지 않는다. 기계 판독과 전수 파일 목록만으로 “누락 없음”을 주장하지 않으며, 미확인 이미지·인용 위치·출처 부족은 감사대장에 남긴다.')
    p('주요 한계는 단일 정류소와 6일의 소표본, 차량 식별 부재, 실제 최초 입주에 대한 사후 증거 부족, 날짜 정밀도 차이, 공표 노선정보의 방향 구분 부족이다. 건축물 사용승인일은 실제 입주일이 아니고, 행정동 세대수는 사업지구 경계와 일치하지 않는다. 따라서 기본서비스 충족 Boolean, 배차 준수율, 누적 피해 세대·일과 인과효과는 산출하지 않았다.')
    sub('재현 및 배포')
    p('원본 검수대장을 보존한 채 결합 결과, 필드별 변경 이력, 원문 날짜 재대조, 보류 목록, 재실행 표와 그림을 V02 폴더로 분리했다. 코드의 고정된 분석 선택과 결과를 미리 넣는 방식은 구분해 문서화했다. 기존의 고정 human_verified=0과 고정 FAIL 요약은 V02 입력으로 사용하지 않는다. 문서 연결 점검은 v3.1-resolution-import로 기록하며 정류소 G1~G6 v2.3과 구별한다.')
    p('배포본은 인증키·개인 행정서류를 제외하고 원본과 변환본의 해시를 기록한다. 과거 키가 포함된 Git 이력은 가져오지 않고 기존의 정제된 비공개 저장소 이력 위에 V02를 추가한다. 연구 표·그림과 별도 감사자료를 함께 제공하여 발표 내용의 근거와 남은 검토 범위를 추적할 수 있도록 했다.')

    h('참고문헌 및 재현 자료',True)
    reference_start=len(d.paragraphs)
    for n,sid in enumerate(['SRC-H001','SRC-H003','SRC-H005','SRC-H007','SRC-H010','SRC-H011','SRC-H025','SRC-H028','SRC-H029'],1):
        x=refs[sid]
        date_text='2011-06-23 (보관 원문 등록일)' if sid=='SRC-H007' else x['published_date']
        if date_text=='Jun-20':date_text='2020년 6월 (대장 기재 월)'
        p(f"[{n}] {x['publisher']}. {x['title']}. {date_text} ({sid}). {x['source_url']}")
    p('[10] 국토해양부. 2011년 6~8월 입주예정 아파트 공개. 2011-06-23 배포, 본문 2쪽. 보관 PDF: 110624(조간)__11년_6~8월_입주예정_아파트_공개(주택정책과).pdf.')
    p('[11] TAGO API 규격 확인자료. docs/observation_plan.md, docs/tago_route_api_findings.md 및 AGENTS.md의 확인 완료 필드 정의. 도착예정시간 단위는 초, 노선 배차간격 단위는 분. 도착정보 원문 활용가이드의 별도 보관 여부는 확인되지 않았다.')
    p('[12] 사용자 검수자료. evidence/candidate_registry.csv; data/source_registry_verified_final_resolution.csv. 197행, 사람 검수 기록 131행. 검수 시각은 2026-09-02 대장 기재값이며 V02 결합 시각과 구분한다.')
    p('[13] 본 연구 재현자료. exports/submission_v02/의 CSV·JSON·그림, docs/submission_audit_v02/의 변경·전수목록·원문대조, scripts/submission/의 V02 분석 코드. 배차 재현 원본은 scratch/calc_headway_v7.py.')
    for para in d.paragraphs[reference_start:]:
        para.paragraph_format.space_after=Pt(4)
        for run in para.runs:run.font.size=Pt(9.5)
    footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('Transport Gap Audit · V02  |  ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for container in [d._element,d.styles._element]:
        for border in container.xpath('.//w:pBdr'):border.getparent().remove(border)
    path=OUT/'교통데이터공모전_논문_V02.docx';d.save(path);print(path)

if __name__=='__main__':build()
