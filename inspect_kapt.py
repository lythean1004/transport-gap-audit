import os
import hashlib
import csv
from pathlib import Path
import openpyxl
from datetime import datetime
from collections import defaultdict

def compute_sha256(filepath: Path) -> str:
    hasher = hashlib.sha256()
    with open(filepath, 'rb') as f:
        for chunk in iter(lambda: f.read(8192), b""):
            hasher.update(chunk)
    return hasher.hexdigest()

def inspect_kapt():
    raw_dir = Path("data/raw/kapt")
    excel_files = list(raw_dir.glob("*.xls*"))
    if not excel_files:
        raise FileNotFoundError("No Excel files found.")
    excel_files.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    target = excel_files[0]
    
    stat = target.stat()
    file_size = stat.st_size
    mtime = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
    sha256 = compute_sha256(target)
    
    wb = openpyxl.load_workbook(target, data_only=True)
    sheet_names = wb.sheetnames
    ws = wb.active
    
    first_5_rows = []
    row_iter = ws.iter_rows(values_only=True)
    for _ in range(5):
        try:
            r = next(row_iter)
            first_5_rows.append(r)
        except StopIteration:
            break
            
    header_idx = -1
    max_strs = -1
    for i, r in enumerate(first_5_rows):
        strs = [str(c) for c in r if c is not None and str(c).strip()]
        if len(strs) > max_strs:
            max_strs = len(strs)
            header_idx = i
            
    header_row = first_5_rows[header_idx]
            
    row_iter = ws.iter_rows(min_row=header_idx+2, values_only=True)
    
    total_rows = 0
    samples = defaultdict(list)
    null_counts = defaultdict(int)
    
    approval_date_cols = []
    ambiguous = []
    for idx, col_name in enumerate(header_row):
        if col_name is None: continue
        cname = str(col_name).strip()
        if "사용승인" in cname or "일" in cname:
            if "사용승인" in cname:
                approval_date_cols.append((idx, cname))
            else:
                ambiguous.append((idx, cname))
                
    approval_date_types = set()
    
    for row in row_iter:
        if not row or not any(row): continue
        total_rows += 1
        for idx, col_name in enumerate(header_row):
            if col_name is None: continue
            val = row[idx] if idx < len(row) else None
            
            if val is None or str(val).strip() == "":
                null_counts[idx] += 1
            else:
                if len(samples[idx]) < 3:
                    samples[idx].append(val)
                    
            if idx in [x[0] for x in approval_date_cols] and val is not None:
                t_name = type(val).__name__
                if hasattr(val, "strftime"):
                    approval_date_types.add(f"datetime (e.g. {val.strftime('%Y-%m-%d')})")
                elif isinstance(val, str):
                    if "-" in val:
                        approval_date_types.add(f"string with hyphens (e.g. {val})")
                    else:
                        approval_date_types.add(f"string without hyphens (e.g. {val})")
                elif isinstance(val, int):
                    approval_date_types.add(f"integer (e.g. {val})")
                else:
                    approval_date_types.add(t_name)
                    
    wb.close()
    
    report_path = Path("docs/kapt_intake_report.md")
    report_path.parent.mkdir(exist_ok=True)
    
    with open(report_path, "w", encoding="utf-8") as f:
        f.write("# K-apt 파일 입고(Intake) 리포트\n\n")
        f.write("## 1. 파일 식별\n")
        f.write(f"- **경로:** {target}\n")
        f.write(f"- **크기(Bytes):** {file_size:,}\n")
        f.write(f"- **mtime(수정일시):** {mtime}\n")
        f.write(f"- **SHA-256:** {sha256}\n\n")
        
        f.write("## 2. 시트 목록\n")
        for s in sheet_names:
            f.write(f"- {s}\n")
        f.write("\n")
        
        f.write("## 3. 주 시트 물리 1~5행 원문 그대로\n")
        f.write("```text\n")
        for i, r in enumerate(first_5_rows):
            clean_r = [str(x) if x is not None else "NULL" for x in r]
            f.write(f"Row {i+1}: {clean_r[:15]} ... (truncated if long)\n")
        f.write("```\n\n")
        
        f.write("## 4. 컬럼 인덱스별 헤더 원문·샘플 3건·null 수\n")
        f.write("| 인덱스 | 헤더 원문 | 샘플 1 | 샘플 2 | 샘플 3 | Null 수 |\n")
        f.write("| --- | --- | --- | --- | --- | --- |\n")
        
        for idx, col_name in enumerate(header_row):
            if col_name is None: continue
            cname = str(col_name).replace('\n', ' ')
            n_count = null_counts[idx]
            samps = samples[idx] + [""]*3
            f.write(f"| {idx} | {cname} | {str(samps[0]).replace('|', '')} | {str(samps[1]).replace('|', '')} | {str(samps[2]).replace('|', '')} | {n_count:,} |\n")
            
        f.write(f"\n## 5. 총 행수(스트리밍 카운트)\n- **{total_rows:,}** 행 (헤더 및 안내문 제외)\n")
        
        f.write("\n## 6. 사용승인일류 컬럼의 실제 저장 타입\n")
        for t in approval_date_types:
            f.write(f"- {t}\n")
            
        f.write("\n## 7. 판단 필요 컬럼 목록\n")
        for idx, cname in ambiguous:
            f.write(f"- [{idx}] {cname}\n")
            
    # Append 1 row to evidence/source_registry.csv
    reg_path = Path("evidence/source_registry.csv")
    with open(reg_path, 'r', encoding='utf-8') as f:
        reader = csv.reader(f)
        header = next(reader)
        
    row_dict = {
        "source_id": "KAPT_EXCEL_LATEST",
        "source_type": "참고자료(주간 갱신)",
        "title": "K-apt 단지 기본정보",
        "publisher": "K-apt",
        "url": "",
        "local_path": str(target).replace('\\', '/'),
        "sha256": sha256,
        "file_size_bytes": file_size,
        "file_mtime": mtime,
        "retrieved_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "notes": "K-apt는 본 엑셀을 참고자료로 명시, 최신 확정값은 K-apt OpenAPI 기준"
    }
    
    with open(reg_path, 'a', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        out_row = [row_dict.get(h, "") for h in header]
        writer.writerow(out_row)

if __name__ == "__main__":
    inspect_kapt()
